from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def extract(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else "",
                    "text": text,
                }
            )

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    return {
        "file": path.name,
        "paragraphs": paragraphs,
        "tables": tables,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
    }


def main() -> None:
    output_path = None
    args = sys.argv[1:]
    if len(args) >= 2 and args[0] == "--output":
        output_path = Path(args[1])
        args = args[2:]
    paths = [Path(arg) for arg in args]
    data = [extract(path) for path in paths]
    if output_path:
        output_path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    else:
        json.dump(data, sys.stdout, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
